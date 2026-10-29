#!/usr/bin/env python3
"""Dump every sheet of the SEO manager's audit workbook, fully, cell by cell."""
import openpyxl

PATH = "/home/z/my-project/upload/SEO Dashboard & Technical Audit _ jadugarrajtilak.com.xlsx"

wb = openpyxl.load_workbook(PATH, data_only=True)
print(f"SHEETS: {wb.sheetnames}\n")

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n{'='*90}\nSHEET: {name}  ({ws.max_row} rows x {ws.max_column} cols)\n{'='*90}")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        values = []
        for cell in row:
            v = cell.value
            if v is None:
                continue
            v = str(v).replace("\n", " ⏎ ")
            values.append(v)
        if not values:
            continue
        # Print row number + joined cells
        print(f"R{row[0].row}: " + " | ".join(values))
