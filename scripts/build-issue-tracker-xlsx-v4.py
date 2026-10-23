#!/usr/bin/env python3
# Jadugar Rajtilak — Issue Tracker & Fix Status workbook (v4 update)
# Updated after: mobile 3-dot menu pre-hydration fix, Contact/About
# redesign, sitemap footer link, professional Leads CRM, agent publish
# execution (blog writer auto-publish + template fallback), live deploy.
import sys, os
XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT = "/home/z/my-project/download/Jadugar_Rajtilak_Issue_Tracker_Fix_Status.xlsx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

wb = Workbook()

# ═══════════════ Sheet 1 — Issue Tracker ═══════════════
ws = wb.active
ws.title = "Issue Tracker"

headers = ["#", "Issue Reported", "Where", "Root Cause (Research)", "Fix Applied", "Status", "Verified On", "Evidence"]
issues = [
    # ── v3 batch (previous session — already live) ──
    (1, "Gallery images appear 3 times (duplicates)", "Gallery page",
     "Neon DB had extra rows gallery.item06/07/09/10 pointing to the SAME uploaded photos; old deployed code rendered 10 frames",
     "Deleted duplicate/override SiteAsset rows; gallery rebuilt: exactly the 7 client photos, one frame each",
     "FIXED", "2026-10-23",
     "Live DOM audit: 7 frames, 7 unique, 0 duplicates"),
    (2, "Home top image - keep it like the old design", "Home (hero)",
     "Repo head contained the 'India's Most Popular' redesign commit; risk of it going live",
     "Confirmed + shipped the ORIGINAL 'Where Reality Bends' edition with the real magician portrait (1024x1365, arch frame)",
     "DONE - OLD STYLE", "2026-10-23",
     "Live hero = real portrait, zero crop"),
    (3, "Home 2nd image - full image should fit in frame", "Home (story section)",
     "Near-square photo (962x944) forced into a square frame -> slight edge crop",
     "Story frame changed to the photo's exact natural ratio (aspect 962/944); rendered via the slot registry",
     "FIXED", "2026-10-23",
     "Live: rendered ratio == natural ratio (fitDelta < 0.02), full photo visible"),
    (4, "Shows page images cut / not fitting in frame", "Shows (services)",
     "Fixed 4:3 landscape frames with portrait photos -> up to 33% of each photo cut off",
     "Every services photo now renders in its own natural-ratio frame from the image slot registry",
     "FIXED", "2026-10-23",
     "Live: all 5 photos fitDelta 0.0000-0.0017 (zero visible crop)"),
    (5, "2 certificates missing - make a proper separate section", "Home + Gallery",
     "Certificates existed only as admin uploads, never placed on the site pages",
     "Dedicated 'Certificates & honours' section on BOTH Home and Gallery with the 2 official scans in gold frames with caption plates",
     "FIXED", "2026-10-23",
     "Live + VLM visual check: Certificate of Recognition + competition poster, framed, titled"),
    (6, "Old certificate showing instead of the admin-uploaded one (LG)", "Certificates",
     "Stale DB override rows pointed at pre-upload assets",
     "v3 one-time migration evicts stale cert/gallery rows on first request; admin-uploaded scans now win",
     "FIXED", "2026-10-23",
     "Live cert section = certificate-intl-2021.webp + certificate-competition-2021.webp (admin scans)"),
    (7, "Site slow / performance dropping (mobile + desktop)", "Whole site",
     "2093KB PNG poster + 2527KB PNG gallery photo + 7 unused stock images; API-bound images blocking render",
     "Recompressed to 241KB/210KB WebP; removed 7 stock files; static assets; lazy loading",
     "FIXED", "2026-10-23",
     "Live perf (both viewports): TTFB 6-29ms, FCP 440-752ms, DCL 390-648ms, JS 70KB"),
    (8, "Broken blog post images", "Blog (2 posts)",
     "2 live posts pointed at deleted /api/media/... URLs",
     "repairFeaturedImages() safety net repoints to identical static webp on first /api/blog hit",
     "FIXED", "2026-10-23",
     "Live blog: 0 broken images, 10 post cards render"),
    # ── v4 batch (this session) ──
    (9, "Mobile: tapping the 3-dot menu does not open sections", "Mobile nav (all pages)",
     "Tap happened BEFORE React hydration finished on slow phones - the button was in the HTML but had no click handler yet, so nothing opened",
     "Pre-hydration fallback: tiny native script in <body> toggles data-jr-menu on <html>; menu panel is always server-rendered; CSS shows it instantly; nav buttons carry data-jr-nav paths and navigate via full page load. React adopts the state on mount and owns the menu afterwards",
     "FIXED", "2026-10-23",
     "Live test with ALL JavaScript blocked (slow-phone simulation): tap opens full-screen menu (390x844) with all 9 sections + booking CTA; About navigates. Post-hydration flow re-verified, zero console errors"),
    (10, "Contact page does not look good - make it better", "Contact page",
     "Plain sand hero + separate status strip felt flat; channel cards stacked in one long list; bottom CTA repeated the hero actions",
     "Redesign: dark plum 'velvet-lobby' hero with gold light, 3 promise chips, one-tap WhatsApp/call CTAs; 4 channel cards in a 2x2 grid; 'Before you write' quick-answers panel linking FAQ; redundant bottom CTA removed",
     "FIXED", "2026-10-23",
     "VLM design review: all core checks PASS (hero premium, grid aligned, form clean). Live mobile + desktop screenshots verified"),
    (11, "About page - make it even better", "About page",
     "Text-only story, thin timeline dots, no numbers band, no pull quote - less premium than the rest of the site",
     "Career stats band (12+ years / 500+ shows / 40+ cities / 4.8 stars, 181 reviews); drop-cap story + gold pull-quote card; milestone timeline in gold-plate cards; touring cities strip + CTA; certificate card retained",
     "FIXED", "2026-10-23",
     "VLM review: PASS on all 6 checks. All admin text-customization keys preserved"),
    (12, "Sitemap not added - add it", "Site-wide",
     "XML sitemap + robots.txt WERE already live (all 11 routes + blog URLs, sitemap referenced in robots) but nothing visible linked to it, so it looked 'missing'",
     "Visible 'Sitemap' link added to the footer bottom bar -> opens /sitemap.xml; XML sitemap auto-includes every published post (agent-published posts appear same day)",
     "PROVIDED + VISIBLE", "2026-10-23",
     "Live: /sitemap.xml 200 (all routes + blog slugs), robots.txt references it, footer link present in live HTML"),
    (13, "Admin Leads CRM not professional - redesign it", "Admin > Leads CRM",
     "Functional but plain: no KPIs, flat cards, and the WhatsApp/email buttons were fake toasts ('Phase 2')",
     "Professional redesign: KPI strip (total / new this week / booked / events within 30 days); kanban with colored column caps, avatars, event countdown badges; polished table; detail dialog with dark header band and REAL one-tap actions - WhatsApp (wa.me with normalized +91 number and prefilled text), call (tel:), email (mailto: with subject/body), copy phone",
     "FIXED", "2026-10-23",
     "Local end-to-end: dialog opens, links verified (wa.me/91982..., tel:..., mailto:...), pipeline move saves + toast, drag-drop retained"),
    (14, "Agents create drafts but never publish/execute", "Admin > AI Agents",
     "Blog Writer saved drafts only by design; publishing required opening Blog Studio for each post; Rank Agent drafts had no execution path",
     "Auto-publish mode (default ON, owner toggle in the panel): Blog Writer publishes the daily post straight to the live site and it enters the sitemap; per-day guard prevents double publishing; 'Pending drafts' board with one-click 'Publish now' + preview; Rank pipeline rows get 'Open platform' buttons (direct listing/signup URLs) + copy-outreach",
     "FIXED", "2026-10-23",
     "Verified run: 'Published live: Magician in Udaipur: The Complete 2026 Booking Guide -> /blog/magician-in-udaipur-2026-booking-guide - already in the sitemap' (writer: llm). Toggle OFF/ON persists. Draft published via button; copy outreach works"),
    (15, "Blog Writer agent shows 'Needs attention'", "Admin > AI Agents",
     "On hosts without LLM credentials the run failed outright (LLM not reachable) -> status stuck on error -> 'Needs attention' chip",
     "Curated template fallback writer (English + Hindi) produces a complete post (6 sections, ~500-580 words, FAQ, internal links) whenever the LLM is unreachable, so every run succeeds; LLM still used when available",
     "FIXED", "2026-10-23",
     "Template unit test: 6 sections / ~582 words (EN) and ~510 words (HI), ALL CHECKS PASS. Agent card now shows 'Active'"),
    (16, "Deploy without any error", "Deployment",
     "Needed a clean lint + build + push + Vercel pipeline",
     "Lint to 0 errors (object-key quoting, JSX nesting, per-render regex, 2 pre-existing rules); production build passed; commit 09560ad pushed to GitHub main; Vercel auto-deploy live",
     "DEPLOYED", "2026-10-23",
     "Live routes all 200; new HTML contains pre-hydration menu script + footer sitemap link; VLM live QA PASS"),
]
last_col = len(headers) + 1
setup_sheet(ws, title="Jadugar Rajtilak - Issue Tracker & Fix Status (v4)", last_col=last_col)
ws.cell(row=3, column=2, value="16 issues tracked - all FIXED / PROVIDED / DEPLOYED and verified. Batch v3: gallery, certificates, image fit, performance. Batch v4: mobile menu, Contact/About redesign, sitemap link, Leads CRM, agent execution, deploy.")
ws.cell(row=3, column=2).font = font_caption()
ws.cell(row=3, column=2).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=last_col)
for c, h in enumerate(headers, start=2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

for i, row in enumerate(issues):
    rn = 5 + i
    for c, v in enumerate(row, start=2):
        ws.cell(row=rn, column=c, value=v)
    style_data_row(ws, row_num=rn, col_start=2, col_end=last_col, row_index=i)
    st = row[5]
    cell = ws.cell(row=rn, column=7)
    if st.startswith(("FIXED", "DONE", "DEPLOYED", "PROVIDED", "VERIFIED")):
        cell.font = Font(name=FONT_NAME, size=11, color=ACCENT_POSITIVE)
        cell.fill = PatternFill('solid', fgColor='E8F5E9')
    ws.cell(row=rn, column=2).alignment = align_number()

ws.freeze_panes = 'C5'
auto_fit_columns(ws, min_width=8, max_width=46, header_row=4, data_start_row=5)
for col in ('C', 'E', 'F', 'I'):
    for r in range(5, 5 + len(issues)):
        ws[f'{col}{r}'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

# ═══════════════ Sheet 2 — Page-by-Page Status ═══════════════
ws2 = wb.create_sheet("Page Status")
h2 = ["Page", "v4 Changes", "Live Status", "Notes"]
rows2 = [
    ("Home", "None this batch (v3 hero/story kept as reverted)", "OK", "Pre-hydration menu fix applies site-wide; footer sitemap link visible"),
    ("About", "Career stats band, drop-cap + pull quote, gold-plate milestones, touring strip + CTA", "REDESIGNED", "All text customization keys preserved for admin edits"),
    ("Shows (Services)", "None this batch", "OK", "Natural-ratio photo frames from v3"),
    ("Packages", "None this batch", "OK", "Card-based design"),
    ("Gallery", "None this batch", "OK", "7 client photos + certificates wall from v3"),
    ("Reviews", "None this batch", "OK", "Text carousel"),
    ("Blog", "Agent posts now auto-publish daily (toggle in AI Agents)", "ENHANCED", "Agent-published post verified live + in sitemap"),
    ("FAQ", "None this batch", "OK", "Contact page links into it"),
    ("Contact", "Velvet-lobby dark hero, promise chips, one-tap CTAs, 2x2 channel grid, quick answers panel", "REDESIGNED", "Form logic unchanged (same validation + admin inbox); VLM PASS"),
    ("Booking", "None this batch", "OK", "Date-check wizard"),
    ("Admin: Leads CRM", "KPI strip, professional kanban + table, real WhatsApp/call/email actions", "REDESIGNED", "One-tap actions verified end-to-end locally"),
    ("Admin: AI Agents", "Auto-publish switch, pending-drafts board, template fallback, Rank execution links", "ENHANCED", "Blog Writer 'Active'; publish verified"),
    ("Footer", "Sitemap link added", "UPDATED", "Opens /sitemap.xml in new tab"),
    ("Mobile navigation", "Pre-hydration menu fallback", "FIXED", "Works with JS blocked (slow-phone simulation) and post-hydration"),
]
last2 = len(h2) + 1
setup_sheet(ws2, title="Page-by-Page Status (v4)", last_col=last2)
for c, h in enumerate(h2, start=2):
    ws2.cell(row=4, column=c, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=last2)
for i, row in enumerate(rows2):
    rn = 5 + i
    for c, v in enumerate(row, start=2):
        ws2.cell(row=rn, column=c, value=v)
    style_data_row(ws2, row_num=rn, col_start=2, col_end=last2, row_index=i)
    st = ws2.cell(row=rn, column=4)
    if st.value in ("REDESIGNED", "ENHANCED", "UPDATED", "FIXED"):
        st.font = Font(name=FONT_NAME, size=11, color=ACCENT_POSITIVE)
auto_fit_columns(ws2, min_width=10, max_width=52, header_row=4, data_start_row=5)
for col in ('C', 'E'):
    for r in range(5, 5 + len(rows2)):
        ws2[f'{col}{r}'].alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

# ═══════════════ Sheet 3 — Performance & Deploy ═══════════════
ws3 = wb.create_sheet("Performance & Deploy")
h3 = ["Metric", "Before v4", "After v4 (live 2026-10-23)", "Result"]
rows3 = [
    ("Mobile 3-dot menu on a slow phone", "Tap in the first seconds = nothing happens (no handler until hydration)", "Opens instantly via native pre-hydration fallback; navigates; React takes over seamlessly", "FIXED + VERIFIED"),
    ("Contact page structure", "Flat hero + stacked list + duplicate bottom CTA", "Premium dark hero, 2x2 channel grid, quick answers; page shorter and sharper", "REDESIGNED"),
    ("About page richness", "Plain text + thin timeline", "Stats band, pull quote, milestone cards, cities strip", "REDESIGNED"),
    ("Sitemap visibility", "XML live but invisible to visitors", "Footer link + auto-includes agent posts", "VISIBLE"),
    ("Leads CRM actions", "Fake buttons (toasts saying 'Phase 2')", "Real wa.me / tel: / mailto: + copy phone, one tap each", "WORKING"),
    ("Agent publishing", "Drafts only, manual publish in Blog Studio", "Auto-publish ON (toggle), per-day guard, one-click publish for old drafts", "EXECUTING"),
    ("Blog Writer status", "'Needs attention' (fails without LLM creds)", "'Active' - template fallback guarantees a daily post", "FIXED"),
    ("ESLint errors", "5 errors (incl. 2 pre-existing)", "0 errors", "CLEAN"),
    ("Production build", "-", "Passed (all 30+ routes compiled)", "PASS"),
    ("Deployment", "c890f7d", "09560ad pushed -> Vercel auto-deploy live", "LIVE"),
    ("All routes HTTP", "-", "200 (11 site routes + sitemap.xml + robots.txt)", "VERIFIED"),
    ("Console/page errors (live)", "-", "None on home/contact/about + admin flows", "CLEAN"),
]
last3 = len(h3) + 1
setup_sheet(ws3, title="Performance & Deployment Summary (v4)", last_col=last3)
for c, h in enumerate(h3, start=2):
    ws3.cell(row=4, column=c, value=h)
style_header_row(ws3, row_num=4, col_start=2, col_end=last3)
for i, row in enumerate(rows3):
    rn = 5 + i
    for c, v in enumerate(row, start=2):
        ws3.cell(row=rn, column=c, value=v)
    style_data_row(ws3, row_num=rn, col_start=2, col_end=last3, row_index=i)
    imp = ws3.cell(row=rn, column=5)
    imp.font = Font(name=FONT_NAME, size=11, color=ACCENT_POSITIVE)
auto_fit_columns(ws3, min_width=10, max_width=48, header_row=4, data_start_row=5)

note_row = 5 + len(rows3) + 2
ws3.cell(row=note_row, column=2, value="How to control the agents: Master Admin -> AI Agents. The 'Auto-publish' switch decides whether the Blog Writer "
        "publishes the daily post live (ON, default) or only saves a draft (OFF). Pending drafts can be published with one click from the same panel. "
        "The Rank Agent never auto-submits to external sites (no-spam rule) - each row has an 'Open platform' button that opens the real listing page, "
        "and outreach drafts can be copied to your email. Existing posts and business data are never modified by agents.")
ws3.cell(row=note_row, column=2).font = font_caption()
ws3.cell(row=note_row, column=2).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
ws3.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=last3)
ws3.row_dimensions[note_row].height = 56

wb.properties.creator = "Z.ai"
wb.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes")
