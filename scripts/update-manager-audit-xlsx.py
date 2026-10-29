#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Update the SEO MANAGER's audit workbook with the 2026-10-29 fixes.

Everything in this run was deployed and live-verified before annotating.
Output: download/SEO_Dashboard_Technical_Audit_UPDATED_2026-10-29.xlsx
"""
from openpyxl import load_workbook

SRC = "/home/z/my-project/upload/SEO Dashboard & Technical Audit _ jadugarrajtilak.com.xlsx"
OUT = "/home/z/my-project/download/SEO_Dashboard_Technical_Audit_UPDATED_2026-10-29.xlsx"

wb = load_workbook(SRC)

# ── 1. Errors Technical & On Page sheet ──────────────────────────
ws = wb["Errors Technical & On Page"]
ERRORS_UPDATE = {
    "T-13": "FIXED 2026-10-29 — six dedicated service pages are LIVE: /services/wedding-magician, /services/corporate-shows, /services/birthday-kids, /services/college-stage, /services/virtual-magic, /services/private-luxury (plus /destination-wedding-magician-udaipur for O-01). Each: 1,100+ words, FAQ, honest pricing, Service+Breadcrumb JSON-LD, audit-supplied metas, crawlable internal links in/out, sitemap entry day one.",
    "T-14": "FIXED 2026-10-29 — /packages expanded 414 → 806 words: honest-pricing context section (four cost drivers explained) + Package FAQ. Booking/contact/legal intros were already shipped 2026-10-29 morning.",
    "T-06": "PARTIALLY FIXED 2026-10-29 — all 13 templated metas rewritten individually (T-08 batch, live). Body differentiation (1,100+ words each) remains a content task: 2–3 posts/week per the calendar.",
    "T-15": "POSITIVE SIGNAL 2026-10-29 — new money pages ship in code-split chunks with full SSR; home bundle unchanged. Field CWV still needs the PSI run + GSC CWV report once access is granted.",
}
for r in range(5, ws.max_row + 1):
    rid = ws.cell(row=r, column=2).value
    if not rid:
        continue
    rid = str(rid).strip()
    if rid in ERRORS_UPDATE:
        ws.cell(row=r, column=4).value = ERRORS_UPDATE[rid]
        print(f"errors sheet: {rid} annotated")

# ── 2. Implementation Roadmap sheet ──────────────────────────────
ws2 = wb["Implementation Roadmap"]
ROADMAP_UPDATE = {
    "T-02a": ("DONE 2026-10-29 (was Week 1)", "Live-verified: sitemap.xml now 43 URLs — 12 core + 7 new money pages + 24 published posts; truthful lastmod per page."),
    "T-06b": ("DONE 2026-10-29 (was Weeks 2–3)", "Live-verified: 23 blog posts now carry internal backlinks to /magician-in-udaipur, /destination-wedding-magician-udaipur and the six service pages — 40 new crawlable internal links, seeded in production DB (idempotent script scripts/seed-internal-backlinks.cjs)."),
    "T-14": ("DONE 2026-10-29 (was Week 3)", "Live-verified: /packages 806 words (honest-pricing context + FAQ)."),
    "R-18": ("DONE 2026-10-29 (was Week 3–4)", "LIVE: /services/wedding-magician — 1,242 words, Service+Breadcrumb JSON-LD, audit meta, inlinks from /services, footer sitewide, homepage grid, 4 wedding-cluster posts."),
    "R-19": ("DONE 2026-10-29 (was Week 4)", "LIVE: /services/corporate-shows — 1,142 words, full schema; inlinks from 5 corporate-cluster posts + /services + footer."),
    "R-20": ("DONE 2026-10-29 (was Week 4–5)", "LIVE: /services/birthday-kids — 1,217 words, full schema; inlinks from 2 birthday posts + /services + footer."),
    "R-22": ("DONE 2026-10-29 (was Weeks 5–8)", "LIVE: /services/college-stage (1,156 w), /services/virtual-magic (1,154 w), /services/private-luxury (1,208 w) — all with Service+Breadcrumb JSON-LD and inlinks."),
    "O-01": ("PAGE LIVE 2026-10-29 — blogs ongoing (was Weeks 5–10)", "LIVE: /destination-wedding-magician-udaipur — 1,267 words, three-day-timeline formats (mehendi / sangeet / villa parlour), local no-travel-overhead positioning, full schema. Destination blog cluster per Nov calendar W1."),
    "T-17a": ("DONE 2026-10-29 (was Week 3)", "LIVE: FAQPage JSON-LD on /faq for the visible Q&A (manager-audit morning batch)."),
}
for r in range(5, ws2.max_row + 1):
    rid = ws2.cell(row=r, column=2).value
    if not rid:
        continue
    rid = str(rid).strip()
    if rid in ROADMAP_UPDATE:
        deadline, validation = ROADMAP_UPDATE[rid]
        ws2.cell(row=r, column=11).value = deadline
        ws2.cell(row=r, column=12).value = validation
        print(f"roadmap: {rid} marked done")

# ── 3. Executive Summary & Scorecard re-score ────────────────────
ws3 = wb["Executive Summary & Scorecard"]
SCORE_UPDATE = {
    "Technical SEO": ("90", "Sitemap coverage complete (43 URLs incl. 8 money pages) + crawlable anchors everywhere; button-nav (T-01) fixed 2026-10-29"),
    "On-page SEO": ("88", "All titles/descriptions in range; per-post metas unique; FAQ schema live"),
    "Content quality": ("70", "10 strong posts + 7 new 1,100+ word money pages; 13 templated bodies still to differentiate (T-06a, content cadence)"),
    "Keyword targeting": ("86", "Every commercial keyword now has a dedicated URL (6 service pages + 2 city/destination pages); cannibalisation resolved (T-05)"),
    "Website architecture": ("86", "Service formats have own URLs; clean flat structure retained"),
    "Internal linking": ("80", "Crawlable nav + footer sitewide + /services section links + homepage grid + 40 seeded post backlinks + money-page link blocks"),
    "OVERALL (average of 8 assessed dimensions)": ("84", "Re-scored 2026-10-29 after Phase-1/2 fixes. Mobile & Performance still Not Assessed (needs GSC/PSI)."),
}
for r in range(1, ws3.max_row + 1):
    dim = ws3.cell(row=r, column=1).value
    if not dim:
        continue
    dim = str(dim).strip()
    if dim in SCORE_UPDATE:
        score, basis = SCORE_UPDATE[dim]
        ws3.cell(row=r, column=2).value = int(score)
        ws3.cell(row=r, column=3).value = basis
        print(f"scorecard: {dim} -> {score}")

wb.save(OUT)
print(f"\nSaved: {OUT}")
